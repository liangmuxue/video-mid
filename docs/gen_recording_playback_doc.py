# -*- coding: utf-8 -*-
"""生成《录像回放技术实现》Word 文档。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parent / "录像回放技术实现.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.35


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 14, 3: 12}
    for run in p.runs:
        set_run_font(run, size=sizes.get(level, 12), bold=True)


def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x22, 0x33, 0x2A)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows, 1):
        for c_i, v in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = str(v)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


# ---------- 正文 ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("极知视频中台 · 录像回放技术实现说明")
set_run_font(r, size=20, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("video-mid Recording Playback Technical Design")
set_run_font(r, size=12)

add_para("版本：V1.0")
add_para(
    "本文档说明 video-mid 中「按天 24 小时时间轴 + 按需拉流」录像回放的完整技术实现，"
    "覆盖数据落盘、后端接口、前端工具包与业务集成。"
)

add_heading_cn("1. 目标与约束", 1)
add_heading_cn("1.1 产品目标", 2)
add_para("选定设备、选定日期后，展示当天 00:00:00 ~ 23:59:59 的录像时间轴。")
add_para("绿色：有录像，可点选 / 拖拽定位；灰色：无录像，不可点、不可拖。")
add_para("播放器底部固定全天时间轴（影视机风格），支持滚轮缩放、框选放大、绿段拖拽选时。")
add_para("按需拉流：只加载当前正在播放的那一个文件；禁止预加载全天录像。")

add_heading_cn("1.2 硬性约束", 2)
add_table(
    ["约束", "实现要点"],
    [
        ("无起止时间选择框", "仅日期；查询窗口固定为当日全天"),
        ("灰区锁死", "hitSegment 未命中则不 seek；灰底 pointer-events: none"),
        ("单路 src", "切换前 unloadVideo；同一时刻仅 1 个 video"),
        ("preload=none", "播放器与 loadAndPlayOne 均强制 none"),
        ("轴在画面下方", ".rdt-player-footer 挂载 DayTimelineBar"),
        ("去掉原生分片进度", "无 controls，避免 xx/300 类控件"),
        ("只读", "工具包不提供录像增删改"),
    ],
)

add_heading_cn("2. 总体架构", 1)
add_para("分层职责如下：")
add_table(
    ["层", "职责"],
    [
        ("落盘", "MainStreamRecordTestTool（或其它录像源）按设备目录写 MP4"),
        ("后端", "按设备 + 时间窗列文件；安全打开文件流"),
        ("工具包", "日期 → 段映射 → 时间轴交互 → 按需设置 video.src"),
        ("业务页", "注入 fetchRecordings / getVideoUrl，弹层或业务端嵌入"),
    ],
)
add_para("数据流示意：")
add_code(
    "业务页 RecordingPlaybackPanel / RecordingDayTimeline\n"
    "  ├─ 元数据：GET .../recordings?from&to  → RecordFileService.list*\n"
    "  └─ 视频字节：GET .../recordings/{id}/{file} → RecordFileService.openFile\n"
    "         ↓\n"
    "  {output-dir}/{deviceId}/yyyyMMdd_HHmmss.mp4"
)

add_heading_cn("3. 数据落盘", 1)
add_heading_cn("3.1 目录结构", 2)
add_para("配置项（application.yml）：")
add_code(
    "testdata:\n"
    "  main-stream-record:\n"
    "    enabled: true\n"
    "    output-dir: data/testdata-records\n"
    "    interval-seconds: 300\n"
    "    clip-seconds: 300          # 与前端 clipSeconds 对齐\n"
    "    ffmpeg-path: ffmpeg\n"
    "    zlm-http-port: 8080"
)
add_para("落盘布局：")
add_code(
    "{output-dir}/\n"
    "  {sanitized_device_id}/\n"
    "    20260915_143000.mp4\n"
    "    20260915_143500.mp4"
)
add_para("文件夹名 = 设备业务 ID（device_id）。文件名 = yyyyMMdd_HHmmss.mp4，表示该段开始时间。")
add_para("若接口无结束时间，前端用 clipSeconds（默认 300）推算绿段结束。")

add_heading_cn("3.2 写入组件", 2)
add_para("类：com.jizhi.videomid.tools.MainStreamRecordTestTool")
add_para(
    "行为概要：定时对主码流用 ffmpeg 截取固定时长片段并写入上述目录。"
    "生产可替换为真实 NVR/边缘录像同步，只要保持「设备目录 + 时间戳文件名」约定，列表逻辑可复用。"
)

add_heading_cn("4. 后端实现", 1)
add_heading_cn("4.1 核心服务 RecordFileService", 2)
add_para("路径：src/main/java/com/jizhi/videomid/record/RecordFileService.java")
add_table(
    ["方法", "作用"],
    [
        ("list(deviceId, from, to)", "扫描设备目录 MP4，按文件名时间过滤、倒序"),
        ("listWithVideoUrls(..., publicBaseUrl)", "在 list 结果上追加 videoUrl"),
        ("openFile(deviceId, fileName)", "安全打开文件为 Resource（防路径穿越）"),
    ],
)
add_para("列表项字段示例：")
add_code(
    "{\n"
    '  "deviceId": "CAM_EAST_01",\n'
    '  "fileName": "20260915_143000.mp4",\n'
    '  "timestamp": "20260915_143000",\n'
    '  "recordTime": "2026-09-15 14:30:00",\n'
    '  "size": 1234567,\n'
    '  "videoUrl": "http://host:8090/api/open/recordings/CAM_EAST_01/20260915_143000.mp4"\n'
    "}"
)
add_para("from / to 支持：yyyyMMdd_HHmmss、yyyy-MM-dd HH:mm:ss、ISO、纯日期。")

add_heading_cn("4.2 接口一览", 2)
add_heading_cn("管理端（需 JWT）", 3)
add_para("控制器：RecordFileController；前缀：/api/recordings")
add_table(
    ["方法", "路径", "说明"],
    [
        ("GET", "/api/recordings?deviceId=&from=&to=", "列表（无 videoUrl）"),
        ("GET", "/api/recordings/{deviceId}/{fileName}", "直出 MP4；video 可用 ?token="),
    ],
)
add_para(
    "鉴权：AuthInterceptor 拦截 /api/**；extractToken 支持 Header Bearer 与 Query token"
    "（因 video 标签无法带 Authorization）。"
)

add_heading_cn("开放接口（免登录）", 3)
add_para("控制器：OpenApiController；前缀：/api/open（WebMvcConfig 排除鉴权）")
add_table(
    ["方法", "路径", "说明"],
    [
        ("GET", "/api/open/devices/{deviceId}/recordings?from=&to=", "列表 + videoUrl"),
        ("GET", "/api/open/recordings/{deviceId}/{fileName}", "公开 MP4"),
    ],
)
add_para("videoUrl 前缀：open-api.public-base-url；为空则按当前请求拼 scheme/host/port。")

add_heading_cn("业务端（需登录）", 3)
add_para("控制器：BizPortalController；前缀：/api/biz")
add_table(
    ["方法", "路径", "说明"],
    [
        (
            "GET",
            "/api/biz/devices/{deviceId}/recordings?from=&to=",
            "带 videoUrl；已停用设备拒绝回放",
        ),
    ],
)
add_para("业务规则：已停用设备灰色可见但不可播（含回放）。")

add_heading_cn("4.3 安全要点", 2)
add_para("openFile：规范化路径，startsWith(设备根目录)，拒绝 ..、非法文件名。")
add_para("文件名须匹配 yyyyMMdd_HHmmss.mp4。")
add_para("开放接口无鉴权，生产应限制内网或另加签名/白名单。")

add_heading_cn("5. 前端工具包", 1)
add_para("目录：web/toolkits/recording-day-timeline/")
add_para("包名：@video-mid/recording-day-timeline；入口：src/index.js")

add_heading_cn("5.1 模块职责", 2)
add_table(
    ["模块", "文件", "职责"],
    [
        ("主组件", "RecordingDayTimeline.vue", "日期、拉列表、播放器壳、底部时间轴"),
        ("时间轴", "DayTimelineBar.vue", "缩放 / 平移 / 绿段拖拽 / 刻度"),
        ("播放器", "OnDemandVideoPlayer.vue", "单路 video preload=none"),
        ("拉流工具", "onDemandPlay.js", "unload / loadAndPlayOne / 禁止预取"),
        ("时间工具", "timeUtils.js", "日界、段合并、命中、选文件"),
        ("缩放工具", "timelineZoom.js", "zoomAt / pan / ticks / 坐标换算"),
    ],
)

add_heading_cn("5.2 RecordingDayTimeline 接口", 2)
add_para("Props：", True)
add_table(
    ["Prop", "类型", "说明"],
    [
        ("deviceId", "string", "必填"),
        ("initialDate", "string", "yyyy-MM-dd，默认今天"),
        ("clipSeconds", "number", "无 endTime 时默认段长，默认 300"),
        ("fetchRecordings", "Function", "必填，(id,{from,to}) => Promise<record[]>"),
        ("getVideoUrl", "Function", "必填，(id,fileName,record) => string"),
        ("showPlayer", "boolean", "是否内置播放器，默认 true"),
    ],
)
add_para("Events：play / seek / loaded / error / date-change")
add_para("Expose：reload / playAt / segments / date")

add_heading_cn("5.3 绿段构建（timeUtils.js）", 2)
add_code(
    "records[]\n"
    "  → 解析 recordTime / timestamp / startTime\n"
    "  → 过滤到当日\n"
    "  → end = endTime 或 start + clipSeconds\n"
    "  → 按 startSec 排序并合并重叠区间 → segments[]"
)
add_para("关键字段：startSec / endSec / records[] / leftPct / widthPct。")
add_para("hitSegment：点是否落在绿段（灰区返回 null）。")
add_para("pickRecordAt：段内按开始时间选中具体文件。")
add_para("earliestSegment：「播放」按钮起点。")

add_heading_cn("5.4 按需拉流（onDemandPlay.js）", 2)
add_code(
    "playAt(daySec)\n"
    "  → hitSegment；未命中则 return\n"
    "  → pickRecordAt → resolveUrl → seekInFile\n"
    "  → playOne(url, seekSeconds)\n"
    "       → 同一文件已加载：只改 currentTime\n"
    "       → 否则：unload → src=url → load → seek → play"
)
add_table(
    ["规则", "说明"],
    [
        ("单路", "新 URL 前必 removeAttribute('src') + load()"),
        ("同文件拖拽", "不重载，避免卡顿"),
        ("预取", "shouldPrefetchNext() 恒为 false"),
        ("连播", "ended 后找下一段绿区再 playAt"),
    ],
)

add_heading_cn("5.5 时间轴交互（DayTimelineBar.vue）", 2)
add_table(
    ["操作", "行为"],
    [
        ("滚轮", "以光标时间为锚点缩放（最小视窗约 10 秒）"),
        ("绿段按住拖拽", "实时 scrub，rAF 合并 seek，触发换流/定位"),
        ("拖到灰区", "不生效，保持上次绿段位置"),
        ("空白拖拽", "框选放大；Alt/Shift+拖拽平移"),
        ("双击 /「全天」/ Esc", "还原 0–24h"),
        ("放大刻度", "由视窗跨度自动切到分、秒级（buildTicks）"),
    ],
)
add_para("播放头仅在绿段上可抓取；灰区播放头不可拖。")

add_heading_cn("5.6 布局结构", 2)
add_code(
    ".rdt-root\n"
    "  .rdt-toolbar          （日期 / 播放 / 刷新）\n"
    "  .rdt-player-shell\n"
    "    .rdt-player-stage   （OnDemandVideoPlayer + 覆盖播放按钮）\n"
    "    .rdt-player-footer  （DayTimelineBar ← 贴底全天轴）"
)
add_para("样式全部 rdt- 前缀 + scoped，避免污染业务页。")

add_heading_cn("6. 业务集成", 1)
add_heading_cn("6.1 RecordingPlaybackPanel", 2)
add_para("路径：web/src/components/RecordingPlaybackPanel.vue")
add_para("封装工具包并提供关闭按钮。默认走管理端 GET /api/recordings + recordingFileUrl（带 token）。")
add_para("可通过 props 覆盖：fetchRecordings、getVideoUrl。")

add_heading_cn("6.2 设备管理弹层", 2)
add_para("路径：web/src/views/DevicesView.vue")
add_para("行操作「录像回放」→ 本页 mask 弹层嵌入 RecordingPlaybackPanel，不新开路由。")
add_para("遮罩关闭：需在遮罩上 pointerdown + pointerup，避免时间轴拖出弹窗外误关。")

add_heading_cn("6.3 业务端", 2)
add_para("路径：web/src/views/BizPortalView.vue")
add_para("目录树选设备 →「录像回放」弹层。")
add_para("注入 fetchBizRecordings（/api/biz/.../recordings）与 videoUrl。")
add_para("已停用设备：playable=false，回放按钮禁用；接口也会拒绝。")

add_heading_cn("6.4 前端 API 封装", 2)
add_para("路径：web/src/api/device.js")
add_table(
    ["函数", "后端"],
    [
        ("fetchRecordings", "GET /api/recordings"),
        ("recordingFileUrl", "/api/recordings/{id}/{file}?token="),
        ("fetchBizRecordings", "GET /api/biz/devices/{id}/recordings"),
    ],
)
add_para("开发环境走 Vite 代理（web/vite.config.js → 后端），避免浏览器直连跨域。")

add_heading_cn("7. 端到端时序", 1)
add_code(
    "用户选日期 / 打开弹层\n"
    "  → dayBounds → from/to = 当日 00:00:00 ~ 23:59:59\n"
    "  → fetchRecordings(deviceId, { from, to })\n"
    "  → buildDaySegments(records, date, 300)\n"
    "  → 渲染灰底 + 绿段；无段则禁用播放\n"
    "\n"
    "用户点「播放」或绿段拖拽/点选\n"
    "  → hitSegment(sec)；灰区忽略\n"
    "  → pickRecordAt → URL + 文件内 seekSeconds\n"
    "  → unload 旧流（若换文件）→ loadAndPlayOne\n"
    "  → timeupdate：playheadSec = 文件日始 + currentTime\n"
    "\n"
    "当前文件 ended\n"
    "  → 查找 playhead 之后的下一段 → playAt（仍按需，不预取）"
)

add_heading_cn("8. 配置清单", 1)
add_table(
    ["配置键", "含义"],
    [
        ("testdata.main-stream-record.enabled", "是否写测试录像"),
        ("testdata.main-stream-record.output-dir", "录像根目录"),
        ("testdata.main-stream-record.clip-seconds", "单段时长（与前端默认 clipSeconds 一致）"),
        ("open-api.public-base-url", "开放 videoUrl 绝对前缀"),
    ],
)

add_heading_cn("9. 关键路径索引", 1)
add_code(
    "后端\n"
    "  record/RecordFileService.java\n"
    "  api/RecordFileController.java\n"
    "  api/OpenApiController.java\n"
    "  api/BizPortalController.java\n"
    "  tools/MainStreamRecordTestTool.java\n"
    "  resources/application.yml\n"
    "\n"
    "工具包\n"
    "  web/toolkits/recording-day-timeline/src/\n"
    "    RecordingDayTimeline.vue / DayTimelineBar.vue\n"
    "    OnDemandVideoPlayer.vue / onDemandPlay.js\n"
    "    timeUtils.js / timelineZoom.js\n"
    "\n"
    "集成\n"
    "  web/src/components/RecordingPlaybackPanel.vue\n"
    "  web/src/views/DevicesView.vue\n"
    "  web/src/views/BizPortalView.vue\n"
    "  web/src/api/device.js"
)

add_heading_cn("10. 本地验证建议", 1)
add_para("1. 确认 output-dir 下有对应 deviceId 的 mp4。")
add_para("2. 管理端设备列表 →「录像回放」：绿段可拖，灰段无效；切换时段 Network 仅出现当前文件请求。")
add_para("3. 业务端：已停用设备不可开回放；已启用可开。")
add_para("4. Demo（可选）：cd web/toolkits/recording-day-timeline && npm run demo → http://localhost:5199/")

add_heading_cn("11. 扩展注意", 1)
add_para("若真实录像有明确结束时间，在列表中返回 endTime，工具包会优先使用，而不再依赖 clipSeconds。")
add_para("若改为对象存储，只需替换 RecordFileService 的 list/open（或改 getVideoUrl 为签名 URL），时间轴与按需播放逻辑可复用。")
add_para("同设备多通道录像时，建议仍按 deviceId 分目录，或在文件名中扩展通道维度并在 list 中过滤。")

doc.save(OUT)
print(f"Wrote: {OUT}")
