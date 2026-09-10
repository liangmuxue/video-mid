# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_run_font(run, size=11, bold=False):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)


def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)


base = "http://8.130.74.232:8090"

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("极知视频中台 · 对外开放接口文档")
set_run_font(r, size=20, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("video-mid Open API")
set_run_font(r, size=12)

add_para("版本：V1.0")
add_para("说明：本组接口位于 /api/open/**，无需登录鉴权，可直接调用。")
add_para("统一响应格式：")
add_code('{\n  "code": 0,\n  "message": "success",\n  "data": ...\n}')
add_para("code = 0 表示成功；非 0 表示失败，message 为错误说明。")

add_heading_cn("一、查询设备", 1)
add_para("根据设备名称 name、设备编码 deviceId 查询设备列表。两个条件均可选，支持模糊匹配；均不传则返回全部设备。")
add_para("请求方式：GET", True)
add_para("接口地址：")
add_code(f"{base}/api/open/devices")
add_para("请求参数（Query）：", True)

table = doc.add_table(rows=4, cols=4)
table.style = "Table Grid"
for i, h in enumerate(["参数名", "类型", "必填", "说明"]):
    table.rows[0].cells[i].text = h
for i, row in enumerate(
    [
        ("name", "string", "否", "设备名称，模糊匹配"),
        ("deviceId", "string", "否", "设备编码，模糊匹配"),
        ("—", "—", "—", "二者都不传时返回全部设备"),
    ],
    1,
):
    for j, v in enumerate(row):
        table.rows[i].cells[j].text = v

add_para("请求示例：")
add_code(
    f"GET {base}/api/open/devices?name=东门\n"
    f"GET {base}/api/open/devices?deviceId=CAM_EAST_01\n"
    f"GET {base}/api/open/devices"
)
add_para("成功响应 data 示例：")
add_code(
    """[
  {
    "id": 1,
    "deviceId": "CAM_EAST_01",
    "name": "东门球机",
    "status": "ON",
    "manufacturer": "宇视",
    "model": "IPC-B系列",
    "address": "小区东门",
    "streamCount": 2
  }
]"""
)

add_heading_cn("二、查询设备码流", 1)
add_para("根据设备 deviceId 查询该设备下全部码流（主码流/子码流）数据。")
add_para("请求方式：GET", True)
add_para("接口地址：")
add_code(f"{base}/api/open/devices/{{deviceId}}/streams")
add_para("路径参数：", True)
table = doc.add_table(rows=2, cols=4)
table.style = "Table Grid"
for i, h in enumerate(["参数名", "类型", "必填", "说明"]):
    table.rows[0].cells[i].text = h
for j, v in enumerate(["deviceId", "string", "是", "设备编码，精确匹配"]):
    table.rows[1].cells[j].text = v

add_para("请求示例：")
add_code(f"GET {base}/api/open/devices/CAM_EAST_01/streams")
add_para("成功响应 data 示例：")
add_code(
    """[
  {
    "id": 10,
    "deviceId": "CAM_EAST_01",
    "streamType": "main",
    "streamUrl": "rtmp://8.130.74.232/live/cam01_main",
    "streamName": "东门-主码流",
    "status": "ON",
    "playCount": 0
  },
  {
    "streamType": "sub",
    "streamUrl": "rtmp://8.130.74.232/live/cam01_sub",
    "streamName": "东门-子码流",
    "status": "ON"
  }
]"""
)
add_para("说明：若设备不存在，返回 code != 0，message 为「设备不存在: xxx」。")

add_heading_cn("三、查询设备历史录像", 1)
add_para(
    "根据设备 deviceId 与时间戳范围，查询该设备历史录制视频，返回每条录像的访问链接 videoUrl。"
)
add_para("请求方式：GET", True)
add_para("接口地址：")
add_code(f"{base}/api/open/devices/{{deviceId}}/recordings")
add_para("路径 / 查询参数：", True)
table = doc.add_table(rows=4, cols=4)
table.style = "Table Grid"
for i, h in enumerate(["参数名", "类型", "必填", "说明"]):
    table.rows[0].cells[i].text = h
for i, row in enumerate(
    [
        ("deviceId", "string", "是", "路径参数，设备编码"),
        ("from", "string", "否", "开始时间戳，含本时刻"),
        ("to", "string", "否", "结束时间戳，含本时刻"),
    ],
    1,
):
    for j, v in enumerate(row):
        table.rows[i].cells[j].text = v

add_para("时间格式支持：", True)
add_para("1）yyyyMMdd_HHmmss，例如 20260910_143000")
add_para("2）yyyy-MM-dd HH:mm:ss，例如 2026-09-10 14:30:00")
add_para("3）yyyy-MM-dd，表示当天 00:00:00")
add_para("from / to 均不传时，返回该设备全部录像文件。")
add_para("请求示例：")
add_code(
    f"GET {base}/api/open/devices/CAM_EAST_01/recordings?from=20260910_000000&to=20260910_235959"
)
add_para("成功响应 data 示例：")
add_code(
    """[
  {
    "deviceId": "CAM_EAST_01",
    "fileName": "20260910_143000.mp4",
    "timestamp": "20260910_143000",
    "recordTime": "2026-09-10T14:30:00",
    "size": 12345678,
    "videoUrl": "http://8.130.74.232:8090/api/open/recordings/CAM_EAST_01/20260910_143000.mp4"
  }
]"""
)
add_para("videoUrl 可直接用浏览器或播放器打开，无需 Token。")

add_heading_cn("四、录像文件直链（附属）", 1)
add_para("请求方式：GET", True)
add_code(f"{base}/api/open/recordings/{{deviceId}}/{{fileName}}")
add_para("返回 Content-Type: video/mp4，可在线播放或下载。此接口同样无需鉴权。")

add_heading_cn("五、错误说明", 1)
table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "场景"
table.rows[0].cells[1].text = "说明"
table.rows[1].cells[0].text = "参数错误"
table.rows[1].cells[1].text = "HTTP 400，code=-1，message 为具体原因"
table.rows[2].cells[0].text = "设备不存在"
table.rows[2].cells[1].text = "查询码流时，message=设备不存在: xxx"
table.rows[3].cells[0].text = "无录像"
table.rows[3].cells[1].text = "返回 code=0，data 为空数组 []"

add_heading_cn("六、配置说明", 1)
add_para("application.yml 中 open-api.public-base-url 用于拼接 videoUrl 绝对地址。")
add_code("open-api:\n  public-base-url: http://8.130.74.232:8090")
add_para("若不配置，则按当前请求的协议/主机/端口自动生成。")
add_para(
    "录像文件目录对应：testdata.main-stream-record.output-dir（默认 data/testdata-records/{deviceId}/yyyyMMdd_HHmmss.mp4）。"
)

out = Path(r"E:\视频\video-mid\docs") / "对外开放接口文档.docx"
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
print("bytes", out.stat().st_size)
