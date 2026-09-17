# -*- coding: utf-8 -*-
"""生成《录像回放进度定位技术说明》Word：流程图 + 技术点（无 matplotlib 依赖）。"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "录像回放进度定位技术说明.docx"


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if name == "Consolas":
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tePr if False else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=10, center=True, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 14, 3: 12}
    for run in p.runs:
        set_run_font(run, size=sizes.get(level, 12), bold=True)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=9, name="Consolas")
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=10)
        shade_cell(table.rows[0].cells[i], "D9E2F3")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_text(table.rows[ri + 1].cells[ci], str(val), size=10, center=False)
    doc.add_paragraph()
    return table


def add_flow_box(doc, text, fill="E8F1FB", decision=False):
    """单列表格作为流程图节点。"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, "FFF4E0" if decision else fill)
    set_cell_text(cell, text, bold=True, size=10, center=True)
    # 固定宽度观感：用段落间距
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
    doc.add_paragraph()
    return table


def add_arrow(doc, label=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    t = f"↓ {label}" if label else "↓"
    run = p.add_run(t)
    set_run_font(run, size=11, bold=True, color=RGBColor(0x2F, 0x6F, 0xED))


def add_branch_row(doc, left_text, right_text, left_fill="F0F0F0", right_fill="E8F1FB"):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(table.rows[0].cells[0], left_text, bold=True, size=9)
    shade_cell(table.rows[0].cells[0], left_fill)
    set_cell_text(table.rows[0].cells[1], "　", size=9)
    set_cell_text(table.rows[0].cells[2], right_text, bold=True, size=9)
    shade_cell(table.rows[0].cells[2], right_fill)
    doc.add_paragraph()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.3

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("录像回放进度定位技术说明")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("拖动 / 点击时间轴进度 · 技术点与流程图")
    set_run_font(r, size=12, color=RGBColor(0x55, 0x55, 0x55))

    add_para(doc, "版本：V1.0")
    add_para(
        doc,
        "适用范围：web/toolkits/recording-day-timeline"
        "（RecordingDayTimeline / DayTimelineBar / onDemandPlay / timeUtils）。",
    )

    # ---- 一 ----
    add_heading_cn(doc, "一、结论摘要", 1)
    add_para(
        doc,
        "拖动播放进度或点击绿区进度时，系统用「当天墙上时钟秒数」定位，再映射到对应的整段 MP4 文件，"
        "最后在该文件内用 HTML5 video.currentTime（秒）跳转。",
        bold=True,
    )
    add_para(doc, "不是 HLS/DASH 切片协议，也不是服务端按帧/GOP 索引推流。")

    add_table(
        doc,
        ["层级", "用什么", "不用什么"],
        [
            ["时间轴 UI", "当天时刻 daySec（0～86399）", "切片序号 / 码流 PTS 直传"],
            ["选内容", "整段 MP4 文件（按开始时间落盘）", "HLS m3u8 / DASH MPD"],
            ["文件内进度", "video.currentTime（秒）", "业务自定义切片索引"],
            ["字节拉取", "浏览器对 MP4 可能发 HTTP Range", "应用层切片清单"],
        ],
    )

    # ---- 二 主流程 ----
    add_heading_cn(doc, "二、主流程图（拖动 / 点击进度）", 1)
    add_para(doc, "下列框图按执行顺序阅读；橙色框为判断节点。")

    add_flow_box(doc, "① 用户拖动绿段滑块，或点击绿区进度", "DCEBFF")
    add_arrow(doc)
    add_flow_box(doc, "② clientX → 当天秒数 daySec\n（clientXToSec，受当前缩放可视窗口影响）")
    add_arrow(doc)
    add_flow_box(doc, "③ 命中绿段？\nhitSegment(segments, daySec)", decision=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("↙ 否（灰区）                          是 ↘")
    set_run_font(run, size=10, bold=True, color=RGBColor(0x66, 0x66, 0x66))

    add_branch_row(
        doc,
        "④a 灰区：忽略\n不 emit seek\n不改播放进度",
        "④b 段内选文件\npickRecordAt(seg, daySec)\n按文件开始时间落入哪条 MP4",
        left_fill="F0F0F0",
        right_fill="E8F1FB",
    )

    add_para(doc, "以下仅「命中绿段」继续：", bold=True)
    add_flow_box(doc, "⑤ 计算文件内偏移\nseekSeconds = daySec − 该文件开始时刻的当天秒数")
    add_arrow(doc)
    add_flow_box(doc, "⑥ 与当前 video.src 是否同一文件？", decision=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("↙ 是（同文件内拖）                    否（跨文件）↘")
    set_run_font(run, size=10, bold=True, color=RGBColor(0x66, 0x66, 0x66))

    add_branch_row(
        doc,
        "⑦a 仅改 video.currentTime\n= seekSeconds\n不换 src、不 unload",
        "⑦b unloadVideo 清 src\n→ loadAndPlayOne(url)\npreload=none 单路加载",
        left_fill="E6F6EC",
        right_fill="FDECEC",
    )

    add_flow_box(doc, "⑧ HTML5 播放器按 currentTime 定位画面\n浏览器可能对 MP4 发 HTTP Range（整文件跳读，非业务切片）")
    add_arrow(doc)
    add_flow_box(doc, "⑨ 播放中 onTimeUpdate：\nplayheadSec = 文件开始秒 + currentTime\n轴上指针与画面同步", "DCEBFF")
    add_arrow(doc)
    add_flow_box(doc, "⑩ ended → 跳下一段绿段起点 playAt\n禁止预取下一段（shouldPrefetchNext=false）")

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tip.add_run(
        "结论：轴用「当天时刻」→ 映射整段 MP4 → 文件内 currentTime（秒）；不是流媒体切片。"
    )
    set_run_font(run, size=10, bold=True, color=RGBColor(0x1A, 0x3A, 0x7A))

    # ---- 三 四层 ----
    add_heading_cn(doc, "三、四层技术模型流程图", 1)

    layers = [
        ("① UI 层 · DayTimelineBar", "E8F1FB",
         "鼠标位置 → 可视窗口比例 → 当天秒数 daySec（0～86399）\n"
         "绿段可 scrub；灰区锁定，不可点、不可拖"),
        ("② 段/文件层 · timeUtils", "E8F1FB",
         "hitSegment：daySec 落在 [startSec, endSec)\n"
         "pickRecordAt：段内按开始时间选中对应整段 MP4（非切片序号）"),
        ("③ 文件内偏移层 · RecordingDayTimeline", "E8F1FB",
         "seekSeconds = daySec − secondsOfDay(文件开始时间)\n"
         "例：文件 14:30:00，点到 14:32:10 → 文件内 seek 130 秒"),
        ("④ 播放器层 · onDemandPlay + HTML5 video", "E8F1FB",
         "同文件：只改 currentTime\n"
         "换文件：unload → 设 src → load → loadedmetadata 后设 currentTime 并 play\n"
         "业务层无 HLS/DASH；Range 由浏览器按需发起"),
    ]
    for i, (title_t, fill, body) in enumerate(layers):
        add_flow_box(doc, f"{title_t}\n{body}", fill)
        if i < len(layers) - 1:
            add_arrow(doc, "下一层")

    # ---- 四 技术点 ----
    add_heading_cn(doc, "四、关键技术点拆解", 1)

    add_heading_cn(doc, "4.1 鼠标位置 → 当天秒数", 2)
    add_para(doc, "文件：DayTimelineBar.vue → scrubAtClientX / clientXToSec")
    add_para(
        doc,
        "根据轨道可视窗口 [viewStart, viewEnd] 与鼠标 clientX，线性换算 daySec。"
        "滚轮缩放、框选放大、平移只改变可视窗口，不改变「一天 = 86400 秒」坐标系。",
    )
    add_code(
        doc,
        "scrubAtClientX(clientX)\n"
        "  sec = clientXToSec(...)\n"
        "  hit = hitSegment(segments, sec)\n"
        "  无 hit → return（灰区）\n"
        "  有 hit → emit('seek', clampedSec, hit)",
    )

    add_heading_cn(doc, "4.2 绿段命中与灰区锁定", 2)
    add_para(doc, "文件：timeUtils.js → hitSegment / buildDaySegments")
    add_para(
        doc,
        "绿段由当天录像列表生成：文件开始时间 → startSec；无 endTime 时用 clipSeconds（默认 300）推 endSec。"
        "灰区无录像，命中失败则不 seek。",
    )
    add_code(doc, "hitSegment: 查找 startSec <= daySec < endSec；找不到返回 null")

    add_heading_cn(doc, "4.3 段内选中哪个 MP4", 2)
    add_para(doc, "文件：timeUtils.js → pickRecordAt")
    add_para(
        doc,
        "按各文件开始时间排序，选取「开始当天秒数 ≤ daySec」的最后一条，即当前时刻所在的那一截 clip。",
    )
    add_para(doc, "落盘：{output-dir}/{deviceId}/yyyyMMdd_HHmmss.mp4（文件名=段开始时间）。")

    add_heading_cn(doc, "4.4 文件内偏移", 2)
    add_para(doc, "文件：RecordingDayTimeline.vue → seekInFile / playAt")
    add_code(
        doc,
        "seekSeconds = daySec - secondsOfDay(文件开始时间)\n"
        "例：14:30:00 的文件，点到 14:32:10 → seekSeconds = 130",
    )

    add_heading_cn(doc, "4.5 同文件 vs 换文件（按需单路）", 2)
    add_para(doc, "文件：onDemandPlay.js → loadAndPlayOne")
    add_table(
        doc,
        ["场景", "行为", "网络"],
        [
            ["同一 MP4 内拖动", "只改 video.currentTime", "通常不重新设 src；浏览器可能 Range 跳读"],
            ["跨到另一 MP4", "unload → 新 src → load → 再设 currentTime", "新发一次文件 GET"],
            ["预取下一段", "禁止", "仅 ended 后再 load 下一段"],
        ],
    )

    add_heading_cn(doc, "4.6 播放头回写", 2)
    add_para(doc, "文件：RecordingDayTimeline.vue → onTimeUpdate")
    add_code(
        doc,
        "playheadSec = secondsOfDay(当前文件开始) + video.currentTime\n"
        "轴指针与画面进度统一在「当天时刻」坐标系",
    )

    # ---- 五 ----
    add_heading_cn(doc, "五、与「切片流」对比", 1)
    add_table(
        doc,
        ["对比项", "本系统", "典型 HLS/DASH 切片"],
        [
            ["内容单元", "整段 MP4（约 clipSeconds，如 5 分钟）", "TS/fMP4 等若干秒一片"],
            ["进度语义", "墙上时钟 + 文件内秒", "片序号或媒体时间线"],
            ["切换成本", "换文件才换 URL", "片列表连续切换"],
            ["服务端协议", "静态文件 / HTTP Range", "m3u8 / MPD 清单"],
        ],
    )

    # ---- 六 ----
    add_heading_cn(doc, "六、关键源码路径", 1)
    add_code(
        doc,
        "DayTimelineBar.vue          # 拖动/点击 → daySec\n"
        "timeUtils.js                # hitSegment / pickRecordAt / buildDaySegments\n"
        "RecordingDayTimeline.vue    # playAt / seekInFile / onTimeUpdate\n"
        "onDemandPlay.js             # unload / loadAndPlayOne / currentTime\n"
        "RecordFileService.java      # 列表与 openFile（不参与秒级 seek 计算）",
    )

    add_heading_cn(doc, "七、一句话备忘", 1)
    add_para(
        doc,
        "轴上用时间（当天秒），内容按文件（整段 MP4），文件内用秒（currentTime）；拖动进度不是切流切片。",
        bold=True,
    )

    doc.save(OUT)
    print("Wrote", OUT)
    print("bytes", OUT.stat().st_size)


if __name__ == "__main__":
    build_doc()
